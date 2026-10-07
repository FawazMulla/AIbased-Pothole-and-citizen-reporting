"""
Builds docs/Project_Documentation.docx from the Markdown files in docs/.
Handles headings, bullets, numbered lists, tables, code blocks, bold/inline code.
Mermaid diagrams are included as code blocks (render them at mermaid.live if needed).

Usage:  python generate_docs_docx.py
"""
import glob
import os
import re

from docx import Document
from docx.shared import Pt

DOCS = os.path.join(os.path.dirname(__file__), "docs")


def add_runs(par, text):
    """Render **bold** and `code` inline markup."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(tok)


def add_code(doc, lines):
    p = doc.add_paragraph()
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)


def convert(doc, md):
    lines, i = md.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-|:]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, val in enumerate(row[: len(rows[0])]):
                    cell = t.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], val)
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        elif ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("# ").strip(), level=min(level, 3))
        elif re.match(r"^\s*[-*] ", ln):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", ln))
        elif ln.startswith(">"):
            add_runs(doc.add_paragraph(), ln.lstrip("> "))
        elif ln.strip():
            add_runs(doc.add_paragraph(), ln)
        i += 1


def main():
    doc = Document()
    doc.add_heading("AI-Based Pothole Detection & Citizen Reporting – Project Documentation", 0)
    for path in sorted(glob.glob(os.path.join(DOCS, "0*.md"))):
        convert(doc, open(path, encoding="utf-8").read())
        doc.add_page_break()
    out = os.path.join(DOCS, "Project_Documentation.docx")
    doc.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
